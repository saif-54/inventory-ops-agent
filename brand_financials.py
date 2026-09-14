import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

# Sample Data
data = [
    ["Classic Black Abaya", 25, 30, 5, 15, 10, 150],
    ["Embroidered Kaftan", 40, 50, 5, 20, 10, 250],
    ["Silk Evening Abaya", 60, 70, 7, 25, 12, 380]
]

columns = [
    "Product_Name", "Fabric_Cost", "Tailoring_Cost", "Packaging", 
    "Expected_CAC", "Shipping", "Selling_Price", "Net_Profit", "Margin_Percentage"
]

# Create Excel workbook and sheet
wb = Workbook()
ws = wb.active
ws.title = "Pricing Model"

# Write headers
ws.append(columns)

# Write data and formulas
# Columns: A:Product, B:Fabric, C:Tailoring, D:Packaging, E:CAC, F:Shipping, G:Selling, H:Net, I:Margin
for i, row in enumerate(data, start=2):
    ws.append(row)
    # Net Profit Formula: Selling - (Fabric + Tailoring + Packaging + CAC + Shipping)
    # H = G - (B + C + D + E + F)
    ws[f"H{i}"] = f"=G{i}-(B{i}+C{i}+D{i}+E{i}+F{i})"
    # Margin Percentage Formula: (Net Profit / Selling)
    ws[f"I{i}"] = f"=H{i}/G{i}"
    ws[f"I{i}"].number_format = '0.00%'

wb.save("agent_workspace/pricing_model.xlsx")
print("Excel sheet 'pricing_model.xlsx' created successfully in 'agent_workspace'.")
